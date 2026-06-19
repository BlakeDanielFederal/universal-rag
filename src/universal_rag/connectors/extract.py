"""File text extraction shared by file-based connectors (SharePoint, local dir).

Routes by extension to a small per-format function. Each is isolated so a corrupt
or unsupported file yields "" (and is skipped) rather than failing the whole sync.
"""

from __future__ import annotations

import io
import os

from universal_rag.connectors.confluence import html_to_text

_TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".json", ".log", ".yaml", ".yml"}
_HTML_EXT = {".html", ".htm"}


def _extract_docx(data: bytes) -> str:
    from docx import Document as Docx

    doc = Docx(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


_EXTRACTORS = {".docx": _extract_docx, ".pptx": _extract_pptx, ".pdf": _extract_pdf}

# Default allowed extensions = everything we can extract.
DEFAULT_EXTENSIONS = frozenset(_EXTRACTORS) | _HTML_EXT | _TEXT_EXT


def extract_text(filename: str, data: bytes, mime: str = "") -> str:
    """Extract plain text from a file's bytes, routing by extension. '' if unsupported."""
    ext = os.path.splitext(filename)[1].lower()
    try:
        if ext in _EXTRACTORS:
            return _EXTRACTORS[ext](data).strip()
        if ext in _HTML_EXT:
            return html_to_text(data.decode("utf-8", errors="replace"))
        if ext in _TEXT_EXT:
            return data.decode("utf-8", errors="replace").strip()
    except Exception:
        return ""  # corrupt/unsupported payload — skip rather than fail the run
    return ""
