from __future__ import annotations

import io
from datetime import datetime

import httpx
import pytest

from universal_rag.connectors.sharepoint import (
    MicrosoftGraphClient,
    acquire_app_token,
    extract_text,
    item_passes,
    item_to_document,
)


# --------------------------------------------------------------------------- #
# OAuth2 token
# --------------------------------------------------------------------------- #
def test_acquire_app_token_posts_client_credentials() -> None:
    captured: dict[str, str] = {}

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/tenant-123/oauth2/v2.0/token"
        body = dict(pair.split("=") for pair in request.content.decode().split("&"))
        captured.update(body)
        return httpx.Response(200, json={"access_token": "TOK", "expires_in": 3600})

    token = acquire_app_token(
        "tenant-123", "client-abc", "secret-xyz", transport=httpx.MockTransport(handler)
    )
    assert token == "TOK"
    assert captured["grant_type"] == "client_credentials"
    assert captured["client_id"] == "client-abc"
    assert captured["scope"] == "https%3A%2F%2Fgraph.microsoft.com%2F.default"


# --------------------------------------------------------------------------- #
# Graph client: auth header, pagination, drive walk
# --------------------------------------------------------------------------- #
def test_client_bearer_auth_and_nextlink_pagination() -> None:
    seen_auth: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen_auth.append(request.headers["Authorization"])
        if request.url.params.get("page") == "2":
            return httpx.Response(200, json={"value": [{"id": "3"}]})
        # first page points to an absolute nextLink
        return httpx.Response(
            200,
            json={
                "value": [{"id": "1"}, {"id": "2"}],
                "@odata.nextLink": "https://graph.microsoft.com/v1.0/sites/s/drives?page=2",
            },
        )

    client = MicrosoftGraphClient(
        "https://graph.microsoft.com/v1.0", "TOK", transport=httpx.MockTransport(handler)
    )
    ids = [d["id"] for d in client._iter("/sites/s/drives")]
    assert ids == ["1", "2", "3"]
    assert seen_auth == ["Bearer TOK", "Bearer TOK"]


def test_delta_paginates_and_returns_delta_link() -> None:
    # page 1 -> @odata.nextLink; page 2 -> @odata.deltaLink (the token for next time)
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.params.get("page") == "2":
            return httpx.Response(
                200,
                json={
                    "value": [{"id": "f2", "name": "b.docx", "file": {}}],
                    "@odata.deltaLink": "https://graph.microsoft.com/v1.0/drives/D/root/delta?token=T2",
                },
            )
        return httpx.Response(
            200,
            json={
                "value": [{"id": "f1", "name": "a.md", "file": {}}],
                "@odata.nextLink": "https://graph.microsoft.com/v1.0/drives/D/root/delta?page=2",
            },
        )

    client = MicrosoftGraphClient(
        "https://graph.microsoft.com/v1.0", "TOK", transport=httpx.MockTransport(handler)
    )
    items, delta_link = client.delta("D")
    assert {i["id"] for i in items} == {"f1", "f2"}
    assert delta_link.endswith("token=T2")


def test_resolve_site_builds_host_path_address() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/v1.0/sites/contoso.sharepoint.com:/sites/Apollo"
        return httpx.Response(200, json={"id": "site-1"})

    client = MicrosoftGraphClient(
        "https://graph.microsoft.com/v1.0", "TOK", transport=httpx.MockTransport(handler)
    )
    assert client.resolve_site("contoso.sharepoint.com:/sites/Apollo")["id"] == "site-1"


# --------------------------------------------------------------------------- #
# Text extraction
# --------------------------------------------------------------------------- #
def test_extract_text_plain_and_html() -> None:
    assert extract_text("notes.md", b"# Title\nhello") == "# Title\nhello"
    html = b"<h1>Roadmap</h1><p>Phase <strong>one</strong> ships Q3.</p>"
    out = extract_text("page.html", html)
    assert "Roadmap" in out and "Phase one ships Q3." in out and "<" not in out


def test_extract_text_docx_real_file() -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Apollo roadmap")
    doc.add_paragraph("Phase one ships the dashboard.")
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("plan.docx", buf.getvalue())
    assert "Apollo roadmap" in text
    assert "Phase one ships the dashboard." in text


def test_extract_text_pptx_real_file() -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly Review"
    buf = io.BytesIO()
    prs.save(buf)
    assert "Quarterly Review" in extract_text("deck.pptx", buf.getvalue())


def test_extract_text_unsupported_and_corrupt_return_empty() -> None:
    assert extract_text("image.png", b"\x89PNG\r\n") == ""  # unsupported extension
    assert extract_text("broken.docx", b"not a real docx") == ""  # corrupt -> swallowed


def test_extract_text_pdf_routes_to_pdf_extractor(monkeypatch: pytest.MonkeyPatch) -> None:
    import universal_rag.connectors.extract as extract

    monkeypatch.setitem(extract._EXTRACTORS, ".pdf", lambda data: "PDF TEXT")
    assert extract_text("report.pdf", b"%PDF-1.7 ...") == "PDF TEXT"


# --------------------------------------------------------------------------- #
# Mapping + filtering helpers
# --------------------------------------------------------------------------- #
def test_item_to_document_maps_fields_and_metadata() -> None:
    item = {
        "id": "01ABC",
        "name": "Apollo Plan.docx",
        "webUrl": "https://contoso.sharepoint.com/sites/Apollo/Plan.docx",
        "size": 4096,
        "lastModifiedDateTime": "2024-03-01T09:30:00Z",
        "file": {
            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        },
        "parentReference": {"path": "/drives/D/root:/Plans"},
        "lastModifiedBy": {"user": {"displayName": "Ada"}},
    }
    doc = item_to_document("apollo-sharepoint", "D", item, "extracted text", kind="sharepoint")
    assert doc.external_id == "D:01ABC"
    assert doc.provider == "sharepoint"
    assert doc.title == "Apollo Plan.docx"
    assert doc.content == "extracted text"
    assert doc.url == item["webUrl"]
    assert doc.updated_at == datetime.fromisoformat("2024-03-01T09:30:00+00:00")
    assert doc.metadata["source_kind"] == "sharepoint"
    assert doc.metadata["drive_id"] == "D"
    assert doc.metadata["author"] == "Ada"
    assert doc.metadata["path"] == "/drives/D/root:/Plans"


def test_item_passes_extension_size_and_recency() -> None:
    allowed = {".docx", ".pdf"}
    base = {"name": "x.docx", "size": 100, "lastModifiedDateTime": "2024-03-10T00:00:00Z"}
    since = datetime.fromisoformat("2024-03-01T00:00:00+00:00")

    assert item_passes(base, since, allowed, max_bytes=1024)
    assert not item_passes({**base, "name": "x.png"}, since, allowed, 1024)  # extension
    assert not item_passes({**base, "size": 99999}, since, allowed, 1024)  # too big
    old = {**base, "lastModifiedDateTime": "2024-02-01T00:00:00Z"}
    assert not item_passes(old, since, allowed, 1024)  # older than cursor
    assert item_passes(base, None, allowed, 1024)  # full scan (no cursor)
